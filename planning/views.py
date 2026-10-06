import calendar
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from datetime import timezone as dt_timezone

from django import forms
from django.contrib import messages
from django.contrib.auth.decorators import permission_required
from django.contrib.auth.mixins import PermissionRequiredMixin
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect
from django.urls import reverse
from django.utils import timezone
from django.views import generic
from django.views.decorators.http import require_POST

from core.fields import WEEKDAY_SHORT

from . import services
from .models import Employee, Holiday, RosterMonth, Shift, WorkFromHomeDay

MONTH_NAMES = ["", "januari", "februari", "maart", "april", "mei", "juni", "juli", "augustus",
               "september", "oktober", "november", "december"]


def _year_month(request):
    today = timezone.localdate()
    try:
        year = int(request.GET.get("jaar", today.year))
        month = int(request.GET.get("maand", today.month))
        date(year, month, 1)
    except ValueError as exc:
        raise Http404 from exc
    return year, month


def _shift_month(year, month, delta):
    month += delta
    while month < 1:
        month, year = month + 12, year - 1
    while month > 12:
        month, year = month - 12, year + 1
    return year, month


def month_data(year, month):
    start, end = services.month_bounds(year, month)
    shifts = {(s.date, s.shift_type): s for s in Shift.objects.filter(date__range=(start, end)).select_related("employee")}
    wfh = defaultdict(list)
    wfh_keys = set()
    for w in WorkFromHomeDay.objects.filter(date__range=(start, end)).select_related("employee"):
        wfh[w.date].append(w)
        wfh_keys.add((w.employee_id, w.date))
    holidays = {h.date: h.name for h in Holiday.objects.filter(date__range=(start, end))}
    absences = defaultdict(set)
    from .models import Absence

    for a in Absence.objects.filter(start_date__lte=end, end_date__gte=start):
        d = max(a.start_date, start)
        while d <= min(a.end_date, end):
            absences[a.employee_id].add(d)
            d += timedelta(days=1)
    return start, end, shifts, wfh, wfh_keys, holidays, absences


class CalendarView(PermissionRequiredMixin, generic.TemplateView):
    permission_required = "planning.view_shift"
    template_name = "planning/calendar.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        year, month = _year_month(self.request)
        start, end, shifts, wfh, wfh_keys, holidays, absences = month_data(year, month)
        today = timezone.localdate()

        weeks = []
        for week in calendar.Calendar(firstweekday=0).monthdatescalendar(year, month):
            days = []
            for day in week[:6]:
                kind = Shift.SATURDAY if day.weekday() == 5 else Shift.EVENING
                days.append({
                    "date": day,
                    "in_month": day.month == month,
                    "is_today": day == today,
                    "shift": shifts.get((day, kind)),
                    "wfh": sorted(wfh.get(day, []), key=lambda w: w.employee.name),
                    "holiday": holidays.get(day),
                })
            weeks.append({"number": week[0].isocalendar()[1], "days": days})

        workdays = [d for d in (start + timedelta(days=i) for i in range((end - start).days + 1)) if d.weekday() < 6]
        employees = Employee.objects.filter(active=True).prefetch_related("departments")
        matrix = []
        counts = defaultdict(lambda: {"wfh": 0, "avond": 0, "zaterdag": 0})
        for e in employees:
            cells = []
            for d in workdays:
                codes = []
                if d in holidays:
                    codes.append(("F", "holiday"))
                elif d in absences[e.id]:
                    codes.append(("A", "absent"))
                elif d.weekday() in e.days_off:
                    codes.append(("V", "off"))
                if (e.id, d) in wfh_keys:
                    codes.append(("T", "wfh"))
                    counts[e.id]["wfh"] += 1
                kind = Shift.SATURDAY if d.weekday() == 5 else Shift.EVENING
                s = shifts.get((d, kind))
                if s and s.employee_id == e.id:
                    codes.append(("Z" if kind == Shift.SATURDAY else "D", "shift"))
                    counts[e.id][kind] += 1
                cells.append({"date": d, "codes": codes, "linked": s is not None and s.employee_id == e.id and s.linked_to_wfh})
            matrix.append({"employee": e, "cells": cells, "counts": counts[e.id]})

        roster = RosterMonth.objects.filter(year=year, month=month).first()
        py, pm = _shift_month(year, month, -1)
        ny, nm = _shift_month(year, month, 1)
        context.update({
            "year": year, "month": month, "month_name": MONTH_NAMES[month],
            "weeks": weeks, "workdays": workdays, "matrix": matrix, "weekday_short": WEEKDAY_SHORT,
            "roster": roster, "prev": {"jaar": py, "maand": pm}, "next": {"jaar": ny, "maand": nm},
            "can_generate": self.request.user.has_perm("planning.generate_roster"),
            "can_edit": self.request.user.has_perm("planning.change_shift"),
            "warnings": [w for w in (roster.notes.splitlines() if roster else []) if w],
        })
        return context


@require_POST
@permission_required("planning.generate_roster", raise_exception=True)
def generate_view(request):
    year, month = int(request.POST["jaar"]), int(request.POST["maand"])
    try:
        roster, result = services.generate_month(year, month, user=request.user)
    except services.RosterLocked as exc:
        messages.error(request, str(exc))
    else:
        messages.success(request, f"Rooster voor {MONTH_NAMES[month]} {year} gegenereerd.")
        for warning in result.warnings:
            messages.warning(request, warning)
    return redirect(f"{reverse('planning:calendar')}?jaar={year}&maand={month}")


@require_POST
@permission_required("planning.generate_roster", raise_exception=True)
def set_status_view(request):
    year, month = int(request.POST["jaar"]), int(request.POST["maand"])
    roster, _ = RosterMonth.objects.get_or_create(year=year, month=month)
    roster.status = request.POST.get("status", "concept")
    roster.save(update_fields=["status", "updated_at"])
    messages.success(request, f"Rooster {MONTH_NAMES[month]} {year} staat nu op {roster.get_status_display().lower()}.")
    return redirect(f"{reverse('planning:calendar')}?jaar={year}&maand={month}")


class ShiftForm(forms.ModelForm):
    class Meta:
        model = Shift
        fields = ["employee", "note"]

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["employee"].queryset = Employee.objects.filter(active=True, in_shift_pool=True)


class ShiftEditView(PermissionRequiredMixin, generic.FormView):
    permission_required = "planning.change_shift"
    template_name = "planning/shift_form.html"
    form_class = ShiftForm

    def dispatch(self, request, *args, **kwargs):
        self.day = date.fromisoformat(kwargs["day"])
        self.kind = Shift.SATURDAY if self.day.weekday() == 5 else Shift.EVENING
        if self.day.weekday() == 6:
            raise Http404
        self.shift = Shift.objects.filter(date=self.day, shift_type=self.kind).first()
        return super().dispatch(request, *args, **kwargs)

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs["instance"] = self.shift or Shift(date=self.day, shift_type=self.kind)
        return kwargs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update({"day": self.day, "shift": self.shift, "kind_label": dict(Shift.TYPE_CHOICES)[self.kind]})
        return context

    def form_valid(self, form):
        if "delete" in self.request.POST and self.shift:
            self.shift.delete()
            messages.success(self.request, "Dienst verwijderd.")
        else:
            shift = form.save(commit=False)
            shift.is_manual = True
            shift.linked_to_wfh = WorkFromHomeDay.objects.filter(employee=shift.employee, date=shift.date).exists()
            shift.save()
            for issue in services.validate_shift(shift):
                messages.warning(self.request, issue)
            messages.success(self.request, f"Dienst op {self.day:%d-%m-%Y} toegewezen aan {shift.employee}.")
        return redirect(f"{reverse('planning:calendar')}?jaar={self.day.year}&maand={self.day.month}")


@require_POST
@permission_required("planning.change_workfromhomeday", raise_exception=True)
def toggle_wfh_view(request):
    employee = get_object_or_404(Employee, pk=request.POST["employee"])
    day = date.fromisoformat(request.POST["date"])
    existing = WorkFromHomeDay.objects.filter(employee=employee, date=day).first()
    if existing:
        existing.delete()
        messages.info(request, f"Thuiswerkdag van {employee} op {day:%d-%m} verwijderd.")
    else:
        WorkFromHomeDay.objects.create(employee=employee, date=day, is_manual=True)
        messages.success(request, f"{employee} werkt op {day:%d-%m} thuis.")
    services.refresh_wfh_links(day, day)
    return redirect(f"{reverse('planning:calendar')}?jaar={day.year}&maand={day.month}#matrix")


def ics_feed(request, token):
    """Persoonlijke agenda-feed (abonneren in Outlook/Google/Apple Agenda)."""
    employee = get_object_or_404(Employee, calendar_token=token, active=True)
    since = timezone.localdate() - timedelta(days=60)
    lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//OPS VET Portaal//NL", "CALSCALE:GREGORIAN",
             f"X-WR-CALNAME:Rooster {employee.name}"]
    stamp = timezone.now().strftime("%Y%m%dT%H%M%SZ")

    def event(uid, summary, day, start=None, end=None):
        lines.extend(["BEGIN:VEVENT", f"UID:{uid}@opsvetportal", f"DTSTAMP:{stamp}", f"SUMMARY:{summary}"])
        if start:
            tz = timezone.get_current_timezone()
            s = timezone.make_aware(datetime.combine(day, start), tz).astimezone(dt_timezone.utc)
            e = timezone.make_aware(datetime.combine(day, end), tz).astimezone(dt_timezone.utc)
            lines.extend([f"DTSTART:{s:%Y%m%dT%H%M%SZ}", f"DTEND:{e:%Y%m%dT%H%M%SZ}"])
        else:
            lines.extend([f"DTSTART;VALUE=DATE:{day:%Y%m%d}", f"DTEND;VALUE=DATE:{day + timedelta(days=1):%Y%m%d}"])
        lines.append("END:VEVENT")

    for s in employee.shifts.filter(date__gte=since):
        if s.shift_type == Shift.SATURDAY:
            event(f"shift-{s.pk}", "Zaterdagdienst", s.date, time(8, 0), time(13, 0))
        else:
            event(f"shift-{s.pk}", "Avonddienst", s.date, time(17, 0), time(21, 0))
    for w in employee.wfh_days.filter(date__gte=since):
        event(f"wfh-{w.pk}", "Thuiswerkdag", w.date)
    lines.append("END:VCALENDAR")
    return HttpResponse("\r\n".join(lines) + "\r\n", content_type="text/calendar; charset=utf-8")


@permission_required("planning.view_shift", raise_exception=True)
def export_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    year, month = _year_month(request)
    start, end, shifts, wfh, wfh_keys, holidays, absences = month_data(year, month)
    wb = Workbook()
    ws = wb.active
    ws.title = f"Thuiswerk {MONTH_NAMES[month]}"
    workdays = [d for d in (start + timedelta(days=i) for i in range((end - start).days + 1)) if d.weekday() < 5]
    ws.append([f"Thuiswerkschema - {MONTH_NAMES[month]} {year}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Naam", "Afdeling / functie"] + [f"{WEEKDAY_SHORT[d.weekday()].capitalize()}\n{d.day}/{d.month}" for d in workdays] + ["Aantal"])
    fills = {"T": PatternFill("solid", fgColor="C6EFCE"), "V": PatternFill("solid", fgColor="D9D9D9")}
    for e in Employee.objects.filter(active=True):
        row = [e.name, e.job_title]
        for d in workdays:
            row.append("T" if (e.id, d) in wfh_keys else ("V" if d.weekday() in e.days_off else ""))
        row.append(sum(1 for v in row[2:] if v == "T"))
        ws.append(row)
        for cell in ws[ws.max_row][2:-1]:
            if cell.value in fills:
                cell.fill = fills[cell.value]
    for cell in ws[2]:
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 30

    ws2 = wb.create_sheet("Avond- en zaterdagdienst")
    ws2.append(["Datum", "Dag", "Type dienst", "Naam", "Gekoppeld aan thuiswerk?"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for (day, kind), s in sorted(shifts.items()):
        ws2.append([day.strftime("%d-%m-%Y"), ["Maandag", "Dinsdag", "Woensdag", "Donderdag", "Vrijdag", "Zaterdag", "Zondag"][day.weekday()],
                    s.get_shift_type_display(), s.employee.name, "n.v.t." if kind == Shift.SATURDAY else ("Ja" if s.linked_to_wfh else "Nee")])
    for col, width in zip("ABCDE", (12, 12, 16, 26, 24)):
        ws2.column_dimensions[col].width = width

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="rooster-{year}-{month:02d}.xlsx"'
    wb.save(response)
    return response


class EmployeeAgendaView(PermissionRequiredMixin, generic.DetailView):
    permission_required = "planning.view_employee"
    model = Employee
    template_name = "planning/employee_detail.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = timezone.localdate()
        e = self.object
        context.update({
            "upcoming_shifts": e.shifts.filter(date__gte=today)[:20],
            "upcoming_wfh": e.wfh_days.filter(date__gte=today)[:20],
            "absences": e.absences.filter(end_date__gte=today),
            "ics_url": self.request.build_absolute_uri(reverse("planning:ics", args=[e.calendar_token])),
        })
        return context
